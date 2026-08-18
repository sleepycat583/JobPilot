from __future__ import annotations

from pathlib import Path


class DocumentExtractionError(ValueError):
    pass


def extract_document_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".txt":
        return _read_text(path)
    if suffix == ".docx":
        return _read_docx(path)
    if suffix == ".pdf":
        return _read_pdf(path)
    raise DocumentExtractionError(f"Unsupported document type: {suffix or 'unknown'}")


def make_text_chunks(text: str, *, size: int = 1_200, overlap: int = 180) -> list[str]:
    compact = "\n".join(line.strip() for line in text.splitlines() if line.strip())
    if not compact:
        return []
    if size <= overlap:
        raise ValueError("Chunk size must exceed overlap")
    chunks: list[str] = []
    start = 0
    while start < len(compact):
        end = min(start + size, len(compact))
        if end < len(compact):
            boundary = max(compact.rfind("\n", start, end), compact.rfind("。", start, end), compact.rfind(".", start, end))
            if boundary > start + size // 2:
                end = boundary + 1
        chunk = compact[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end >= len(compact):
            break
        start = max(end - overlap, start + 1)
    return chunks


def _read_text(path: Path) -> str:
    data = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise DocumentExtractionError("TXT file encoding is not supported")


def _read_docx(path: Path) -> str:
    from docx import Document

    try:
        document = Document(path)
    except Exception as exc:  # python-docx has no stable public exception hierarchy
        raise DocumentExtractionError("DOCX file could not be read") from exc
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                paragraphs.append(" | ".join(cells))
    return "\n".join(paragraphs)


def _read_pdf(path: Path) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(path)
        if reader.is_encrypted:
            raise DocumentExtractionError("Encrypted PDF is not supported")
        return "\n".join((page.extract_text() or "").strip() for page in reader.pages)
    except DocumentExtractionError:
        raise
    except Exception as exc:  # pypdf parser errors vary by version
        raise DocumentExtractionError("PDF file could not be read") from exc
